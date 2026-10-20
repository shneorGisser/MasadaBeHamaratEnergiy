"""Word (.docx) export of a report snapshot — the same immutable snapshot the HTML/PDF render.

Formatting follows the course's report conventions: Hebrew right-to-left, David 11 pt, line
spacing 1.15, 8 pt after each paragraph, sources marked (model / simulation / measurement), and
an appendix with MATLAB code that redraws every curve from the report's own tables (so a
student can rebuild the figures in the format their instructor wants)."""

from __future__ import annotations

import io
import math
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "David"
SIZE = Pt(11)
SIM = RGBColor(0x1F, 0x6F, 0xB2)
MEAS = RGBColor(0xA2, 0x46, 0x2E)
THEORY = RGBColor(0x4B, 0x6B, 0x2F)
WARN = RGBColor(0x7A, 0x4A, 0x07)
METHOD_HE = {"consistent": "עקבי", "manual": "נוסחת החוברת", "analytic": "פאזורים (סימולציה בלבד)"}


def _fmt(v: Any, sig: int = 4) -> str:
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "כן" if v else "לא"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if not math.isfinite(v):
            return "—"
        if v == 0:
            return "0"
        a = abs(v)
        if a >= 1e6 or a < 1e-3:
            return f"{v:.{sig - 1}e}"
        digits = max(0, sig - 1 - math.floor(math.log10(a)))
        return f"{v:.{min(digits, 6)}f}"
    return str(v)


def _pct(v: Any) -> str:
    return "—" if v is None else f"{100 * v:.2f}%"


def _strip(html: str) -> str:
    """Plain text of the small HTML snippets in report_info (sub/sup become _x / ^x)."""
    t = re.sub(r"<sub>(.*?)</sub>", r"_\1", html)
    t = re.sub(r"<sup>(.*?)</sup>", r"^\1", t)
    return re.sub(r"<[^>]+>", "", t)


# --------------------------------------------------------------------------- low-level RTL


def _rtl_paragraph(p: Any) -> None:
    ppr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.insert(0, bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _style_run(r: Any, *, bold: bool = False, size: Any = None, color: Any = None) -> None:
    r.font.name = FONT
    r.font.size = size or SIZE
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), FONT)
    rtl = OxmlElement("w:rtl")
    rpr.append(rtl)
    cs = OxmlElement("w:szCs")
    cs.set(qn("w:val"), str(int((size or SIZE).pt * 2)))
    rpr.append(cs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))


class _Doc:
    def __init__(self) -> None:
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = SIZE
        pf = st.paragraph_format
        pf.line_spacing = 1.15
        pf.space_after = Pt(8)
        pf.space_before = Pt(0)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(k), FONT)
        # the whole section right-to-left (page order, columns)
        sect = self.d.sections[0]._sectPr
        sect.append(OxmlElement("w:bidi"))

    def p(
        self,
        text: str = "",
        *,
        bold: bool = False,
        size: Any = None,
        color: Any = None,
        ltr: bool = False,
    ) -> Any:
        para = self.d.add_paragraph()
        if not ltr:
            _rtl_paragraph(para)
        if text:
            r = para.add_run(text)
            _style_run(r, bold=bold, size=size, color=color)
            if ltr:
                r._r.get_or_add_rPr().remove(r._r.get_or_add_rPr().find(qn("w:rtl")))
        return para

    def heading(self, text: str, level: int) -> None:
        self.p(text, bold=True, size=Pt({0: 18, 1: 14, 2: 12}[level]))

    def bullets(self, items: list[str]) -> None:
        """Real Word bullets: in a right-to-left paragraph the bullet sits on the right."""
        for it in items:
            para = self.d.add_paragraph(style="List Bullet")
            _rtl_paragraph(para)
            _style_run(para.add_run(it))

    def note(self, text: str, color: Any = WARN) -> None:
        self.p(text, color=color, size=Pt(10))

    def table(self, header: list[str], rows: list[list[str]], color: Any = None) -> None:
        if not header:
            return
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        tblpr.append(OxmlElement("w:bidiVisual"))
        # the header row repeats on every page the table continues on
        trpr = t.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
        size = Pt(9.5) if len(header) <= 7 else Pt(8)
        for j, h in enumerate(header):
            self._cell(t.cell(0, j), h, bold=True, size=size)
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                self._cell(t.cell(i, j), v, color=color, size=size)
        self.d.add_paragraph()

    def _cell(
        self, cell: Any, text: str, *, bold: bool = False, color: Any = None, size: Any = None
    ) -> None:
        para = cell.paragraphs[0]
        _rtl_paragraph(para)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        r = para.add_run(text)
        _style_run(r, bold=bold, size=size or Pt(9.5), color=color)

    def code(self, text: str) -> None:
        for line in text.splitlines():
            para = self.d.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            r = para.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9)

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.d.save(buf)
        return buf.getvalue()


# --------------------------------------------------------------------------- MATLAB appendix


def _mat(xs: list[Any]) -> str:
    return (
        "["
        + " ".join("NaN" if not isinstance(v, (int, float)) else repr(float(v)) for v in xs)
        + "]"
    )


def _ident(s: str) -> str:
    x = re.sub(r"\W", "_", s, flags=re.ASCII).strip("_") or "x"
    return x if x[0].isalpha() else "v_" + x


def matlab_appendix(snapshot: dict[str, Any]) -> str:
    """MATLAB script that redraws the report's curves from its own tables."""
    out = [
        f"% {snapshot.get('title', '')} — curves rebuilt from the report's tables",
        "% Generated by the Virtual Energy Lab; data = the report snapshot, nothing recomputed.",
        "close all; clc;",
    ]
    fig = 0
    for sec in snapshot.get("sections", []):
        run = sec.get("run") or {}
        res = run.get("result") or {}
        tables = {t.get("id"): t for t in res.get("tables", [])}
        meas = (sec.get("measurements") or {}).get("rows", [])
        for curve in res.get("curves", []):
            tb = tables.get(curve.get("table") or "readings") or (res.get("tables") or [None])[0]
            if not tb:
                continue
            xk, yk = curve["x"], curve["y"]
            labels = {
                c["key"]: f"{c['label_he']}" + (f" [{c['unit']}]" if c.get("unit") else "")
                for c in tb["columns"]
            }
            fig += 1
            rows = tb["rows"]
            out += [
                "",
                f"%% {sec.get('title_he', '')} — {curve.get('title_he', '')}",
                f"figure({fig}); hold on; grid on;",
                f"x_sim = {_mat([r.get(xk) for r in rows])};",
                f"y_sim = {_mat([r.get(yk) for r in rows])};",
                "plot(x_sim, y_sim, '-o', 'DisplayName', 'Simulation');",
            ]
            if meas and any(isinstance(r.get(xk), (int, float)) for r in meas):
                out += [
                    f"x_meas = {_mat([r.get(xk) for r in meas])};",
                    f"y_meas = {_mat([r.get(yk) for r in meas])};",
                    "plot(x_meas, y_meas, 's', 'DisplayName', 'Measurement');",
                ]
            out += [
                f"xlabel('{_ident(xk)}'); ylabel('{_ident(yk)}');  % {labels.get(xk, xk)} / "
                f"{labels.get(yk, yk)}",
                f"title('{_ident(curve.get('id', 'curve'))}'); legend('Location', 'best');",
            ]
    if fig == 0:
        out.append("% (no curves in this report)")
    return "\n".join(out)


# --------------------------------------------------------------------------- document


def build_docx(snapshot: dict[str, Any]) -> bytes:
    s = snapshot
    d = _Doc()
    d.heading(s["title"], 0)
    d.note(
        f"פרויקט: {s['project']['name']} · ניסוי: {s['experiment']['id']} "
        f"v{s['experiment']['version']} · מודל: {s['model_version']} · הופק: "
        f"{s['generated_at'][:16].replace('T', ' ')} UTC",
        color=None,
    )
    d.p("מקרא צבעים: ירוק = מודל / תיאוריה · כחול = סימולציה · חום-אדום = מדידה", size=Pt(9.5))

    info = s["info"]
    d.heading("1. מטרת הניסוי", 1)
    d.bullets(info["goals"])

    d.heading(f"2. {info['equipment_title']}", 1)
    spec = s["spec"]
    if spec.get("provenance") == "assumed":
        d.note(f"פרמטרי {s['equipment_he']} מוערכים ואינם נתוני שלט: {spec.get('provenance_note')}")
    elif spec.get("provenance") == "fitted":
        d.note(f"פרמטרי {s['equipment_he']} הותאמו לנתונים: {spec.get('provenance_note')}")
    d.table(
        ["גודל", "ערך"],
        [[_strip(h), _fmt(v)] for h, v in info["spec_table"]],
        color=THEORY,
    )
    d.p(_strip(info["model_html"]))
    d.p("לא ממודל: " + info["not_modelled"])

    for i, sec in enumerate(s["sections"]):
        d.heading(f"{i + 3}. {sec['title_he']} ({sec['manual_ref']})", 1)
        run = sec.get("run")
        if run:
            d.heading("תנאי עבודה וסימולציה", 2)
            d.note(
                f"ריצה {run['id'][:8]} · {run.get('label') or ''} · סטטוס {run['status']} · "
                f"hash {run['input_hash'][:12]}",
                color=SIM,
            )
            case = run.get("case") or {}
            if case:
                d.note(
                    "תנאי הניסוי: " + " · ".join(f"{k} = {v}" for k, v in case.items()), color=None
                )
            if run.get("error"):
                d.note("הסימולציה נכשלה: " + str(run["error"].get("message")))
            for w in run.get("warnings") or []:
                d.note(str(w.get("message")))
            res = run.get("result") or {}
            for tb in res.get("tables", []):
                d.p(tb["title_he"], bold=True)
                cols = tb["columns"]
                d.table(
                    [c["label_he"] + (f" [{c['unit']}]" if c.get("unit") else "") for c in cols],
                    [[_fmt(r.get(c["key"])) for c in cols] for r in tb["rows"]],
                    color=SIM,
                )
            if res.get("parameters"):
                d.p("חישובים מתוך תוצאות הסימולציה", bold=True)
                d.table(
                    ["גודל", "ערך", "נוסחה", "שיטה"],
                    [
                        [
                            p["label_he"],
                            (
                                f"{_fmt(p['value'], 5)} {p['unit']}"
                                if p["value"] is not None
                                else f"לא מוגדר — {p.get('undefined_reason')}"
                            ),
                            p.get("formula", ""),
                            METHOD_HE.get(p.get("method", ""), p.get("method", "")),
                        ]
                        for p in res["parameters"]
                    ],
                )
        m = sec.get("measurements")
        if m:
            d.heading("מדידות במעבדה", 2)
            st = m["set"]
            d.note(
                f"סט: {st['label']} · תאריך: {st.get('lab_date') or '—'} · מכשירים: "
                f"{st.get('instruments') or '—'}"
                + (
                    " · תנאים: "
                    + " · ".join(f"{k} = {_fmt(v)}" for k, v in st["conditions"].items())
                    if st.get("conditions")
                    else ""
                ),
                color=MEAS,
            )
            if st.get("notes"):
                d.note(st["notes"], color=None)
            cols = m["columns"]
            d.table(
                ["#"]
                + [c["label_he"] + (f" [{c['unit']}]" if c.get("unit") else "") for c in cols],
                [
                    [str(r.get("row_label") or r.get("point"))]
                    + [_fmt(r.get(c["key"])) for c in cols]
                    for r in m["rows"]
                ],
                color=MEAS,
            )
            for w in (m.get("analysis") or {}).get("warnings", []):
                d.note("בדיקת סבירות: " + str(w.get("message")))
            if m.get("missing"):
                d.note(f"חסרים {len(m['missing'])} ערכים בטבלת המדידות.")
        c = sec.get("comparison")
        if c:
            d.heading("השוואה: מדידה מול סימולציה", 2)
            d.note(
                f"שגיאה יחסית מרבית: {_pct(c['summary']['max_abs_rel_error'])} · חציונית: "
                f"{_pct(c['summary']['median_abs_rel_error'])} · שורות ללא התאמה: "
                f"{c['summary']['unmatched_rows']}",
                color=None,
            )
            rows: list[list[str]] = []
            for pr in c["pairs"]:
                if pr["simulated_row"] is None:
                    rows.append(
                        [str(pr["measured_row"] + 1), str(pr["unmatched_reason"]), "", "", ""]
                    )
                    continue
                for cell in pr["cells"]:
                    if cell["measured"] is None:
                        continue
                    rows.append(
                        [
                            str(pr["measured_row"] + 1),
                            cell["quantity"] + (f" [{cell['unit']}]" if cell["unit"] else ""),
                            _fmt(cell["measured"]),
                            _fmt(cell["simulated"]),
                            _pct(cell["rel_error"])
                            if cell["rel_error"] is not None
                            else (cell.get("undefined_reason") or "—"),
                        ]
                    )
            d.table(["שורה", "גודל", "נמדד", "מסומלץ", "שגיאה יחסית"], rows)
            diag = c.get("diagnosis") or []
            if diag:
                d.p("הסברים אפשריים לפערים (השערות שזוהו אוטומטית — לבדוק מול הרישום)", bold=True)
                d.bullets([f["message"] for f in diag])
            if c.get("parameters"):
                d.table(
                    ["פרמטר", "מודל", "מחולץ מסימולציה", "מחולץ ממדידה", "שגיאה יחסית"],
                    [
                        [
                            p["label_he"],
                            _fmt(p["theory"], 5),
                            _fmt(p["simulated"], 5),
                            _fmt(p["measured"], 5),
                            _pct(p["rel_error"])
                            if p["rel_error"] is not None
                            else (p.get("undefined_reason") or "—"),
                        ]
                        for p in c["parameters"]
                    ],
                )
            d.note("ההפרש כולל יחד שגיאת מודל ושגיאת מדידה.", color=None)

    n = len(s["sections"])
    d.heading(f"{n + 3}. דיון", 1)
    if s.get("discussion"):
        for line in str(s["discussion"]).splitlines():
            d.p(line)
    else:
        d.note("הסטודנט לא כתב דיון. המערכת אינה מנסחת הסבר לפערים במקום הסטודנט.")
    d.heading(f"{n + 4}. מגבלות ונתונים חסרים", 1)
    lim = []
    if spec.get("provenance") == "assumed":
        lim.append(info["assumed_limitation"])
    for sec in s["sections"]:
        if not sec.get("run"):
            lim.append(f"{sec['title_he']}: אין סימולציה בדו״ח.")
        if not sec.get("measurements"):
            lim.append(f"{sec['title_he']}: אין מדידות בדו״ח.")
    lim.append(info["general_limitation"])
    d.bullets(lim)

    d.heading("נספח א׳ — קוד MATLAB לשרטוט העקומים", 1)
    d.p("הקוד משרטט מחדש כל עקום מתוך טבלאות הדו״ח (סימולציה ומדידה), ללא חישוב נוסף.")
    d.code(matlab_appendix(s))
    return d.bytes()
