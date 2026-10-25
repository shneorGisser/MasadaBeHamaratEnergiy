"""Word export of a report: Hebrew RTL, David 11, line spacing 1.15, 8 pt after, the report's
tables, the automatic explanations, and a MATLAB appendix that redraws the curves."""

from __future__ import annotations

import io

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient


def test_docx_export(client: TestClient) -> None:
    pid = client.post(
        "/api/projects", json={"name": "דו״ח", "experiment_id": "synchronous"}
    ).json()["id"]
    sid = client.post(
        f"/api/projects/{pid}/datasets/synchronous_2025/import",
        json={"study_kind": "short_circuit"},
    ).json()["sets"][0]["id"]
    run = client.post(f"/api/measurement-sets/{sid}/simulate").json()
    rep = client.post(
        f"/api/projects/{pid}/reports",
        json={
            "title": "ניסוי 3 — קצרים",
            "sections": [{"study_kind": "short_circuit", "run_id": run["id"], "set_id": sid}],
            "student_discussion": "הזרמים גבוהים פי 1.8 מהמודל.",
        },
    )
    assert rep.status_code == 201, rep.text
    r = client.get(f"/api/reports/{rep.json()['id']}/docx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    doc = Document(io.BytesIO(r.content))
    normal = doc.styles["Normal"]
    assert normal.font.name == "David" and normal.font.size.pt == 11
    assert normal.paragraph_format.line_spacing == 1.15
    assert normal.paragraph_format.space_after.pt == 8
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ניסוי 3 — קצרים" in text and "הזרמים גבוהים פי 1.8" in text
    assert "הסברים אפשריים לפערים" in text  # the comparison's diagnosis is in the report
    assert "נספח א׳ — קוד MATLAB" in text and "plot(x_sim, y_sim" in text
    assert "x_meas" in text  # measured points are drawn too
    # right-to-left paragraphs and tables
    rtl = [
        p
        for p in doc.paragraphs
        if p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None
    ]
    assert len(rtl) > 10
    assert doc.tables and doc.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    assert len(doc.tables) >= 4  # spec, simulation, measurements, comparison
    # every curve of the section is a picture, captioned
    pics = doc.inline_shapes
    assert len(pics) >= 1
    assert any(p.text.startswith("איור 3.1") for p in doc.paragraphs)


def test_long_traces_are_folded_and_charted(client: TestClient) -> None:
    """A transient run's trace (thousands of rows) is not printed as a 100-page table: the
    report shows the first rows and the chart."""
    pid = client.post(
        "/api/projects", json={"name": "ט", "experiment_id": "transformer3ph"}
    ).json()["id"]
    run = client.post(f"/api/projects/{pid}/runs", json={"study_kind": "inrush", "case": {}}).json()
    rep = client.post(
        f"/api/projects/{pid}/reports",
        json={"title": "זרם כניסה", "sections": [{"study_kind": "inrush", "run_id": run["id"]}]},
    )
    assert rep.status_code == 201, rep.text
    doc = Document(io.BytesIO(client.get(f"/api/reports/{rep.json()['id']}/docx").content))
    assert max(len(t.rows) for t in doc.tables) <= 16 + 1
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "מוצגות 15 מתוך" in text
    assert len(doc.inline_shapes) == 2  # line currents, limb fluxes


def test_html_report_has_charts_and_folded_tables_and_data_zip(client: TestClient) -> None:
    import csv
    import zipfile

    pid = client.post(
        "/api/projects", json={"name": "ט", "experiment_id": "transformer3ph"}
    ).json()["id"]
    run = client.post(f"/api/projects/{pid}/runs", json={"study_kind": "inrush", "case": {}}).json()
    rep = client.post(
        f"/api/projects/{pid}/reports",
        json={"title": "זרם כניסה", "sections": [{"study_kind": "inrush", "run_id": run["id"]}]},
    ).json()
    html = client.get(f"/api/reports/{rep['id']}/html").text
    assert html.count('<img src="data:image/png;base64,') == 2
    assert "מוצגות 15 מתוך" in html
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/reports/{rep['id']}/data.zip").content))
    names = z.namelist()
    assert "README.txt" in names and "s1_inrush_trace.csv" in names
    rows = list(csv.reader(io.StringIO(z.read("s1_inrush_trace.csv").decode())))
    trace = next(t for t in run["result"]["tables"] if t["id"] == "trace")
    assert rows[0] == [c["key"] for c in trace["columns"]] and len(rows) - 1 == len(trace["rows"])
    # the MATLAB appendix is thinned for print and points to the full CSV
    doc = Document(io.BytesIO(client.get(f"/api/reports/{rep['id']}/docx").content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "readtable('s1_inrush_trace.csv')" in text
